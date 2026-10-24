"""Rights Decay's data: read resolution times, fit hazard curves, read the FCA's UK figures, and
draw the reliability chart (build plan R3). Used by scripts/rights_backtest.py.

The NYC data is US housing repairs standing in for how often a UK company sorts a problem before
you chase it; the FCA's "upheld" share is the UK anchor for how many closures actually pay.
"""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import yaml
from PIL import Image, ImageDraw, ImageFont

from tempo_core.delay_model import Bin, Tally, prior_strength
from tempo_core.hazard import Hazard

KNOTS = (0, 1, 2, 3, 5, 7, 10, 14, 21, 30, 45, 60, 75, 90, 120, 150, 180)
FETCHED = datetime(2026, 10, 24)  # open requests are censored at the download date
MIN_DETAIL_N = 50
SOURCE = "NYC 311 housing requests 2025"
NOTE = (
    "US housing repair requests as a stand-in for how often a UK company sorts a problem before "
    "you chase it. Kaplan-Meier on all of 2025; scored out of sample in docs/rights_calibration.md."
)

Curve = tuple[float, ...]  # P(closed by each knot)


@dataclass(frozen=True)
class Case:
    created: datetime
    days: float  # until closed, or until FETCHED if still open (censored)
    closed: bool
    kind: str
    detail: str


def read_cases(path: Path) -> list[Case]:
    cases = []
    with path.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            created = datetime.fromisoformat(row["created_date"])
            closed_at = datetime.fromisoformat(row["closed_date"]) if row["closed_date"] else None
            closed = row["status"] == "Closed" and closed_at is not None
            days = ((closed_at if closed and closed_at else FETCHED) - created).total_seconds()
            if days >= 0:
                cases.append(
                    Case(created, days / 86400, closed, row["complaint_type"], row["descriptor"])
                )
    return cases


def km(cases: Sequence[Case]) -> Curve:
    """Kaplan-Meier P(closed by each knot); closings come before censorings at equal times."""
    ordered = sorted((c.days, not c.closed) for c in cases)
    at_risk, survive, out, i = len(ordered), 1.0, [0.0], 0
    for knot in KNOTS[1:]:
        while i < len(ordered) and ordered[i][0] <= knot:
            if not ordered[i][1]:
                survive *= 1 - 1 / at_risk
            at_risk, i = at_risk - 1, i + 1
        out.append(1 - survive)
    return tuple(out)


def at(curve: Curve, day: float) -> float:
    """The same straight lines between knots as Hazard.p_resolved_by."""
    if day <= 0:
        return 0.0
    for d1, p1, d2, p2 in zip(KNOTS, curve, KNOTS[1:], curve[1:], strict=False):
        if day <= d2:
            return p1 + (p2 - p1) * (day - d1) / (d2 - d1)
    return curve[-1]


@dataclass(frozen=True)
class Fitted:
    all: Curve
    kinds: dict[str, tuple[Curve, int]]
    details: dict[tuple[str, str], Curve]  # each shrunk toward its kind

    def curve(self, case: Case) -> Curve:
        if (case.kind, case.detail) in self.details:
            return self.details[(case.kind, case.detail)]
        return self.kinds.get(case.kind, (self.all, 0))[0]


def fit(cases: Sequence[Case]) -> Fitted:
    by_kind: dict[str, list[Case]] = defaultdict(list)
    by_detail: dict[tuple[str, str], list[Case]] = defaultdict(list)
    for c in cases:
        by_kind[c.kind].append(c)
        by_detail[(c.kind, c.detail)].append(c)
    kinds = {k: (km(cs), len(cs)) for k, cs in by_kind.items()}
    at14 = KNOTS.index(14)
    details = {}
    for kind, (parent, _) in kinds.items():
        groups = {
            g: (km(cs), len(cs))
            for g, cs in by_detail.items()
            if g[0] == kind and len(cs) >= MIN_DETAIL_N
        }
        tallies = [Tally(round(cv[at14] * n), n) for cv, n in groups.values()]
        strength = prior_strength(tallies, parent[at14])  # method of moments at 14 days
        for g, (cv, n) in groups.items():
            details[g] = tuple(
                (p * n + q * strength) / (n + strength) for p, q in zip(cv, parent, strict=True)
            )
    return Fitted(km(cases), kinds, details)


def hazard(curve: Curve, *, id: str, n: int, share: float = 1.0, share_source: str = "") -> Hazard:
    rounded = [Decimal(f"{p:.4f}") for p in curve]
    knots = tuple((d, max(rounded[: i + 1])) for i, d in enumerate(KNOTS))  # keep it monotone
    return Hazard(
        id=id, source=SOURCE, n=n, demo_values=True, note=NOTE, resolved_by=knots,
        pays_share=Decimal(f"{share:.4f}"), pays_share_source=share_source,
    )  # fmt: skip


def slug(text: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


def write_hazard(folder: Path, h: Hazard) -> None:
    body: dict[str, object] = {"id": h.id, "source": h.source, "n": h.n, "demo_values": True}
    body["note"] = h.note
    if h.pays_share < 1:
        body |= {"pays_share": str(h.pays_share), "pays_share_source": h.pays_share_source}
    body["resolved_by"] = [[d, str(p)] for d, p in h.resolved_by]
    text = yaml.safe_dump(
        body, sort_keys=False, allow_unicode=True, width=100, default_flow_style=None
    )
    header = "# Fitted by scripts/rights_backtest.py (make rights-backtest). Do not edit by hand.\n"
    (folder / f"{h.id}.yaml").write_text(header + text, encoding="utf-8")


@dataclass(frozen=True)
class UkFigures:
    curves: list[Hazard]
    upheld_share: float  # weighted by complaints closed, banking and credit cards
    firms: int


def fca(path: Path, top: int = 15) -> UkFigures | None:
    """Per firm: closed within 3 days and within 8 weeks, and the share upheld (FCA, H1 2026)."""
    try:
        import openpyxl  # type: ignore[import-untyped]  # only `make rights-backtest` brings it
    except ImportError:
        print("openpyxl isn't installed: skipping the FCA figures (make rights-backtest adds it)")
        return None
    book = openpyxl.load_workbook(path, read_only=True)
    after = next(s for s in book.sheetnames if s.startswith("Percentage after"))
    sheets = {
        "closed": "Closed", "d3": "Percentage within 3 days", "w8": after,
        "upheld": "Percentage upheld",
    }  # fmt: skip
    values: dict[str, dict[str, float]] = {}
    for key, sheet in sheets.items():
        rows = list(book[sheet].iter_rows(values_only=True))
        col = rows[0].index("Banking and credit cards")
        values[key] = {str(r[0]): float(r[col]) for r in rows[1:] if r[0] and r[col] is not None}
    firms = [f for f in values["closed"] if all(f in values[k] for k in ("d3", "w8", "upheld"))]
    total = sum(values["closed"][f] for f in firms)
    share = sum(values["upheld"][f] * values["closed"][f] for f in firms) / total
    curves = []
    for firm in sorted(firms, key=lambda f: -values["closed"][f])[:top]:
        d3, w8, up = (values[k][firm] for k in ("d3", "w8", "upheld"))
        knots = (
            (0, Decimal(0)),
            (3, Decimal(f"{d3:.4f}")),
            (56, Decimal(f"{min(1.0, d3 + w8):.4f}")),
        )
        curves.append(Hazard(
            id=f"uk.fca.{slug(firm)}", source="FCA firm-level complaints, H1 2026",
            n=int(values["closed"][firm]), demo_values=True, resolved_by=knots,
            pays_share=Decimal(f"{up:.4f}"),
            pays_share_source="this firm's upheld share, FCA H1 2026",
            note="Banking and card complaints this firm closed within 3 days and within 8 weeks. "
            "Coarse by construction: two points, then flat.",
        ))  # fmt: skip
    return UkFigures(curves, share, len(firms))


def draw(path: Path, series: dict[str, list[Bin]], label: str) -> None:
    size, margin = 560, 70
    image = Image.new("RGB", (size, size), "white")
    pen = ImageDraw.Draw(image)
    font, small = ImageFont.load_default(size=14), ImageFont.load_default(size=12)
    plot = size - 2 * margin

    def xy(px: float, py: float) -> tuple[float, float]:
        return margin + px * plot, size - margin - py * plot

    pen.rectangle([xy(0, 1), xy(1, 0)], outline="#334155")
    pen.line([xy(0, 0), xy(1, 1)], fill="#94a3b8", width=2)
    for tick in (0, 0.25, 0.5, 0.75, 1):
        x, y = xy(tick, 0)
        pen.text((x - 10, y + 6), f"{tick:.2f}", fill="#334155", font=small)
        x, y = xy(0, tick)
        pen.text((x - 40, y - 7), f"{tick:.2f}", fill="#334155", font=small)
    colours = ("#0f766e", "#b45309", "#7c3aed")
    for i, (name, points) in enumerate(series.items()):
        centres = [xy(b.predicted, b.observed) for b in points]
        if len(centres) > 1:
            pen.line(centres, fill=colours[i], width=2)
        for b, (cx, cy) in zip(points, centres, strict=True):
            r = 3 + min(b.runs, 40_000) ** 0.5 / 20
            pen.ellipse([cx - r, cy - r, cx + r, cy + r], fill=colours[i])
        y = margin + 12 + 18 * i
        pen.ellipse([margin + 10, y + 3, margin + 20, y + 13], fill=colours[i])
        pen.text((margin + 26, y), f"closed {name}", fill=colours[i], font=font)
    pen.text((margin, 20), "Rights Decay: how often a problem is closed", fill="#0f172a", font=font)
    pen.text((margin, 40), label, fill="#334155", font=small)
    pen.text((size / 2 - 60, size - 30), "predicted probability", fill="#334155", font=small)
    pen.text((8, margin - 22), "observed", fill="#334155", font=small)
    image.save(path)
